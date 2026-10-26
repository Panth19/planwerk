"""Excel in, Excel out.

Lawyers trust their spreadsheet. So the export is not a dump of numbers: it is
a working Excel model with live formulas and named parameters, and the test
suite recalculates it in headless LibreOffice to prove it gives the same result
as the engine (tests/test_excel_parity.py).

The import reads an employee list with whatever German or English headers the
client used. Known synonyms are mapped deterministically; unknown headers are
reported back so a human (or the header-mapping assistant, which only ever sees
the header names, never the rows) can resolve them.
"""

from __future__ import annotations

import csv
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from .engine import Result
from .model import Case, Employee, PointsScheme, Scheme

EUR = '#,##0 "€";-#,##0 "€"'
EUR_CENT = '#,##0.00 "€"'
DATE = "DD.MM.YYYY"

TENURE_CODE = {"months": 1, "full_years": 2, "half_up": 3}
FORMULA_CODE = {"factor": 1, "divisor": 2}

# --------------------------------------------------------------------------- #
# Export
# --------------------------------------------------------------------------- #

# (name, label, value-getter, number format)
PARAMS: list[tuple[str, str, Any, str | None]] = [
    ("Stichtag", "Stichtag", lambda s, p: s.stichtag, DATE),
    ("Formel", "Formel (1 = Faktor, 2 = Divisor)", lambda s, p: FORMULA_CODE[s.formula], "0"),
    ("Faktor", "Faktor", lambda s, p: s.factor, "0.000"),
    ("Divisor", "Divisor", lambda s, p: s.divisor, "0.0"),
    ("BZRundung", "BZ-Rundung (1 = Monate, 2 = volle Jahre, 3 = ab 6 Monaten)", lambda s, p: TENURE_CODE[s.tenure_rounding], "0"),
    ("MinBZ", "Mindest-BZ in Monaten", lambda s, p: s.min_tenure_months, "0"),
    ("Sockel", "Sockelbetrag", lambda s, p: s.floor_amount, EUR),
    ("Deckel", "Höchstbetrag (0 = keiner)", lambda s, p: s.cap_amount, EUR),
    ("RenteAlter", "Rentennah ab Alter (0 = keine Regel)", lambda s, p: s.retirement_age, "0"),
    ("RenteMonate", "Rentennah: max. Monatsgehälter", lambda s, p: s.retirement_cap_months, "0.0"),
    ("ZKind", "Zuschlag je Kind", lambda s, p: s.child_bonus, EUR),
    ("ZSB", "Zuschlag Schwerbehinderung", lambda s, p: s.disability_bonus, EUR),
    ("Volumen", "Sozialplanvolumen (0 = keins)", lambda s, p: s.budget, EUR),
    ("Kuerzung", "Anteilige Kürzung (1 = ja)", lambda s, p: 1 if s.budget_mode == "pro_rata" else 0, "0"),
    ("Schritt", "Rundung auf (EUR)", lambda s, p: s.rounding_step, "0"),
    ("PAlter", "Punkte je Lebensjahr", lambda s, p: p.per_year_age, "0"),
    ("PBZ", "Punkte je Dienstjahr (1.-10.)", lambda s, p: p.per_year_tenure, "0"),
    ("PBZ10", "Punkte je Dienstjahr (ab 11.)", lambda s, p: p.per_year_tenure_after_10, "0"),
    ("PKind", "Punkte je Kind", lambda s, p: p.per_child, "0"),
    ("PVerh", "Punkte verheiratet", lambda s, p: p.married, "0"),
    ("PSB", "Punkte Schwerbehinderung", lambda s, p: p.disabled, "0"),
]

# Columns of Anlage 1: header, width, number format, formula template ({r} = row)
COLUMNS: list[tuple[str, int, str | None, str | None]] = [
    ("Personal-Nr.", 12, None, None),                                   # A
    ("Name", 22, None, None),                                           # B
    ("Geburtsdatum", 13, DATE, None),                                   # C
    ("Eintritt", 13, DATE, None),                                       # D
    ("Brutto/Monat", 13, EUR, None),                                    # E
    ("Kinder", 8, "0", None),                                           # F
    ("Verh. (1/0)", 10, "0", None),                                     # G
    ("SB (1/0)", 9, "0", None),                                         # H
    ("Alter", 8, "0", 'DATEDIF(C{r},Stichtag,"y")'),                    # I
    ("BZ Monate", 10, "0", 'IF(D{r}>Stichtag,0,DATEDIF(D{r},Stichtag,"m"))'),  # J
    ("BZ Jahre", 10, "0.00", "CHOOSE(BZRundung,J{r}/12,INT(J{r}/12),INT((J{r}+6)/12))"),  # K
    ("Berechtigt", 10, "0", "IF(J{r}>=MinBZ,1,0)"),                      # L
    ("Formelbetrag", 14, EUR_CENT, "IF(Formel=1,K{r}*E{r}*Faktor,I{r}*K{r}*E{r}/Divisor)"),  # M
    ("Nach Sockel", 14, EUR_CENT, "IF(AND(Sockel>0,M{r}<Sockel),Sockel,M{r})"),  # N
    ("Nach Deckel", 14, EUR_CENT, "IF(AND(Deckel>0,N{r}>Deckel),Deckel,N{r})"),  # O
    ("Grundbetrag", 14, EUR_CENT,
     "IF(AND(RenteAlter>0,I{r}>=RenteAlter,O{r}>RenteMonate*E{r}),RenteMonate*E{r},O{r})"),  # P
    ("Zuschläge", 12, EUR, "F{r}*ZKind+IF(H{r}=1,ZSB,0)"),             # Q
    ("Vor Kürzung", 14, EUR_CENT, "IF(L{r}=1,P{r}+Q{r},0)"),            # R
    ("Abfindung", 14, EUR,
     # ROUND(...,6) mirrors engine.clean(): float noise must not flip a rounding decision
     "IF(Quote<1,ROUNDDOWN(ROUND(ROUND(R{r}*Quote,6)/Schritt,6),0),"
     "ROUND(ROUND(ROUND(R{r}*Quote,6)/Schritt,6),0))*Schritt"),  # S
    ("Sozialpunkte", 12, "0",
     "I{r}*PAlter+MIN(INT(J{r}/12),10)*PBZ+MAX(INT(J{r}/12)-10,0)*PBZ10"
     "+F{r}*PKind+IF(G{r}=1,PVerh,0)+IF(H{r}=1,PSB,0)"),                 # T
]
AMOUNT_COL = "S"
PREQUOTA_COL = "R"

INK = "1B2230"
STAMP = "3B47B8"
PAPER = "EEF0EA"
thin = Side(style="thin", color="C9CDC2")


def export_workbook(case: Case, path: str | Path, result: Result | None = None,
                    points: PointsScheme | None = None) -> Path:
    """Write a self-calculating Excel model of the case."""
    s, p = case.scheme, points or PointsScheme()
    wb = Workbook()

    # --- Parameter sheet --------------------------------------------------- #
    ws = wb.active
    ws.title = "Parameter"
    ws["A1"] = case.title
    ws["A1"].font = Font(bold=True, size=14, color=INK)
    ws["A2"] = "Gelbe Felder sind Eingaben. Alle Beträge in Anlage 1 rechnen live mit."
    ws["A2"].font = Font(italic=True, color="5B6470")
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 16
    first = 4
    for i, (name, label, get, fmt) in enumerate(PARAMS):
        r = first + i
        ws.cell(r, 1, label)
        c = ws.cell(r, 2, get(s, p))
        c.fill = PatternFill("solid", fgColor="FFF4C2")
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        if fmt:
            c.number_format = fmt
        wb.defined_names[name] = DefinedName(name, attr_text=f"Parameter!$B${r}")

    n = len(case.employees)
    last = n + 1
    r = first + len(PARAMS) + 1
    ws.cell(r, 1, "Summe vor Kürzung").font = Font(bold=True)
    ws.cell(r, 2, f"=SUM('Anlage 1'!{PREQUOTA_COL}2:{PREQUOTA_COL}{last})").number_format = EUR
    wb.defined_names["SummeBrutto"] = DefinedName("SummeBrutto", attr_text=f"Parameter!$B${r}")
    ws.cell(r + 1, 1, "Kürzungsquote").font = Font(bold=True)
    ws.cell(r + 1, 2, "=IF(AND(Kuerzung=1,Volumen>0,SummeBrutto>Volumen),Volumen/SummeBrutto,1)"
            ).number_format = "0.0000%"
    wb.defined_names["Quote"] = DefinedName("Quote", attr_text=f"Parameter!$B${r + 1}")
    ws.cell(r + 2, 1, "Summe Abfindungen").font = Font(bold=True)
    ws.cell(r + 2, 2, f"=SUM('Anlage 1'!{AMOUNT_COL}2:{AMOUNT_COL}{last})").number_format = EUR
    ws.cell(r + 3, 1, "Abstand zum Volumen").font = Font(bold=True)
    ws.cell(r + 3, 2, f"=IF(Volumen>0,Volumen-B{r + 2},0)").number_format = EUR

    # --- Anlage 1 ---------------------------------------------------------- #
    a = wb.create_sheet("Anlage 1")
    for ci, (head, width, fmt, _) in enumerate(COLUMNS, start=1):
        c = a.cell(1, ci, head)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=INK)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        a.column_dimensions[get_column_letter(ci)].width = width
    a.row_dimensions[1].height = 30
    a.freeze_panes = "C2"

    for ri, e in enumerate(case.employees, start=2):
        values = [e.id, e.name, e.birth_date, e.entry_date, e.monthly_gross,
                  e.children, int(e.married), int(e.disabled)]
        for ci, (head, _, fmt, formula) in enumerate(COLUMNS, start=1):
            v = values[ci - 1] if ci <= len(values) else f"={formula.format(r=ri)}"
            c = a.cell(ri, ci, v)
            if fmt:
                c.number_format = fmt
            if head == "Abfindung":
                c.font = Font(bold=True, color=STAMP)

    tr = last + 1
    a.cell(tr, 2, "Summe").font = Font(bold=True)
    for col in ("Q", "R", "S"):
        c = a[f"{col}{tr}"]
        c.value = f"=SUM({col}2:{col}{last})"
        c.number_format = EUR
        c.font = Font(bold=True)
    a.auto_filter.ref = f"A1:T{last}"

    # --- Notes ------------------------------------------------------------- #
    h = wb.create_sheet("Hinweise")
    notes = [
        "Erzeugt mit Planwerk. Die Formeln entsprechen der Engine 1:1 (Paritätstest in der Test-Suite).",
        "Synthetische Demodaten." if case.synthetic else "Vertraulich: enthält Mandantendaten.",
        "Rechenweg: Formelbetrag → Sockelbetrag → Höchstbetrag → rentennahe Begrenzung → + Zuschläge → Kürzungsquote → Rundung.",
        "Bei anteiliger Kürzung wird abgerundet, damit das Volumen nicht überschritten wird.",
        "Keine Rechtsberatung. Ergebnisse vor Verwendung anwaltlich prüfen.",
    ]
    h.column_dimensions["A"].width = 120
    for i, t in enumerate(notes, start=1):
        h.cell(i, 1, t)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def read_amounts(path: str | Path) -> dict[str, float]:
    """Read computed amounts from a recalculated workbook (cached values)."""
    wb = load_workbook(path, data_only=True)
    a = wb["Anlage 1"]
    out: dict[str, float] = {}
    col = next(i for i, c in enumerate(COLUMNS, start=1) if c[0] == "Abfindung")
    for row in a.iter_rows(min_row=2, values_only=True):
        if row[0] is None or row[1] == "Summe":
            continue
        out[str(row[0])] = float(row[col - 1])
    return out


# --------------------------------------------------------------------------- #
# Import
# --------------------------------------------------------------------------- #

FIELD_SYNONYMS: dict[str, list[str]] = {
    "id": ["personalnummer", "personalnr", "persnr", "pnr", "id", "employeeid", "staffno", "nr"],
    "name": ["name", "mitarbeiter", "mitarbeiterin", "vornamename", "employee", "fullname", "nachnamevorname"],
    "birth_date": ["geburtsdatum", "gebdatum", "geboren", "birthdate", "dateofbirth", "dob", "gebam"],
    "entry_date": ["eintritt", "eintrittsdatum", "beschaeftigtseit", "betriebszugehoerigkeitab",
                   "entrydate", "startdate", "hiredate", "seit"],
    "monthly_gross": ["brutto", "bruttomonatsgehalt", "monatsbrutto", "gehalt", "bruttogehalt",
                      "monthlygross", "salary", "grosssalary", "bruttomonat"],
    "children": ["kinder", "unterhaltspflichtigekinder", "anzahlkinder", "children", "kids"],
    "married": ["verheiratet", "familienstand", "verh", "married", "ehe"],
    "disabled": ["schwerbehindert", "sb", "schwerbehinderung", "gdb50", "disabled", "disability"],
}
REQUIRED = ("birth_date", "entry_date", "monthly_gross")


def _norm(s: str) -> str:
    s = re.sub(r"\(.*?\)", "", s)  # "Verh. (1/0)" -> "Verh."
    s = s.strip().lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    s = unicodedata.normalize("NFKD", s)
    return re.sub(r"[^a-z0-9]", "", s)


@dataclass
class ImportReport:
    employees: list[Employee]
    mapping: dict[str, str]  # field -> original header
    unmapped_headers: list[str]
    missing_fields: list[str]
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.missing_fields


def map_headers(headers: Iterable[str]) -> tuple[dict[str, str], list[str]]:
    mapping: dict[str, str] = {}
    unmapped: list[str] = []
    for h in headers:
        if h is None or str(h).strip() == "":
            continue
        key = _norm(str(h))
        hit = next((f for f, syn in FIELD_SYNONYMS.items() if key in syn and f not in mapping), None)
        if hit:
            mapping[hit] = str(h)
        else:
            unmapped.append(str(h))
    return mapping, unmapped


def _rows(path: Path) -> tuple[list[str], list[list[Any]]]:
    if path.suffix.lower() in {".csv", ".txt"}:
        text = path.read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=";,\t")
        rows = list(csv.reader(text.splitlines(), dialect))
        return rows[0], rows[1:]
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["Anlage 1"] if "Anlage 1" in wb.sheetnames else wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    header_idx = next(i for i, r in enumerate(rows) if sum(v is not None for v in r) >= 3)
    return [str(v) if v is not None else "" for v in rows[header_idx]], rows[header_idx + 1:]


def _date(v: Any) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Datum nicht lesbar: {s!r}")


def _money(v: Any) -> float:
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("€", "").replace("EUR", "").strip()
    if "," in s:  # German format 4.250,50
        s = s.replace(".", "").replace(",", ".")
    return float(s)


def read_employees(path: str | Path, overrides: dict[str, str] | None = None) -> ImportReport:
    """Read an employee list. ``overrides`` maps field -> header for unknown columns."""
    path = Path(path)
    headers, rows = _rows(path)
    mapping, unmapped = map_headers(headers)
    for f, h in (overrides or {}).items():
        mapping[f] = h
        if h in unmapped:
            unmapped.remove(h)
    missing = [f for f in REQUIRED if f not in mapping]
    report = ImportReport([], mapping, unmapped, missing)
    if missing:
        return report

    idx = {f: headers.index(h) for f, h in mapping.items()}
    for n, row in enumerate(rows, start=2):
        if row is None or all(v in (None, "") for v in row):
            continue
        if str(row[idx.get("name", 0)] or "").strip().lower() in {"summe", "total"}:
            continue
        try:
            def get(f, default=None):
                return row[idx[f]] if f in idx and row[idx[f]] not in (None, "") else default
            employees_id = str(get("id") or f"MA-{n - 1:03d}")
            report.employees.append(Employee.from_json({
                "id": employees_id,
                "name": str(get("name") or employees_id),
                "birth_date": _date(get("birth_date")).isoformat(),
                "entry_date": _date(get("entry_date")).isoformat(),
                "monthly_gross": _money(get("monthly_gross")),
                "children": int(float(get("children", 0))),
                "married": get("married", False),
                "disabled": get("disabled", False),
            }))
        except (ValueError, TypeError) as exc:
            report.warnings.append(f"Zeile {n}: {exc}")
    return report


def case_from_file(path: str | Path, scheme: Scheme | None = None) -> Case:
    rep = read_employees(path)
    if not rep.ok:
        raise ValueError(f"Pflichtspalten fehlen: {rep.missing_fields}; nicht zugeordnet: {rep.unmapped_headers}")
    return Case(title=Path(path).stem, employees=rep.employees, scheme=scheme or Scheme(), synthetic=False)
