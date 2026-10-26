import json

from openpyxl import load_workbook

from planwerk.cli import _parse_scenarios, main


def test_calc_prints_a_german_summary(capsys):
    assert main(["calc"]) == 0
    out = capsys.readouterr().out
    assert "Berechtigt" in out and "135 von 140" in out and "€" in out


def test_calc_json_and_set(capsys):
    assert main(["calc", "--json", "--set", "factor=0.5", "--set", "budget_mode=pro_rata"]) == 0
    data = json.loads(capsys.readouterr().out)
    assert len(data["lines"]) == 140


def test_fit(capsys):
    assert main(["fit", "--budget", "6500000"]) == 0
    fit = json.loads(capsys.readouterr().out)
    assert fit["feasible"] and fit["parameter"] == "factor"


def test_export(tmp_path, capsys):
    out = tmp_path / "a.xlsx"
    assert main(["export", "-o", str(out)]) == 0
    assert load_workbook(out).sheetnames == ["Parameter", "Anlage 1", "Hinweise"]


def test_errors_are_one_line_not_a_traceback(capsys):
    assert main(["calc", "--set", "faktor=0.8"]) == 2
    assert main(["calc", "does-not-exist.json"]) == 2
    assert main(["calc", "--set", "nonsense"]) == 2
    err = capsys.readouterr().err
    assert err.count("planwerk:") == 3 and "Traceback" not in err


def test_scenario_syntax():
    assert _parse_scenarios(["A=", "B=factor=0.6,cap_amount=120000,formula=divisor"]) == {
        "A": {},
        "B": {"factor": 0.6, "cap_amount": 120000, "formula": "divisor"},
    }
